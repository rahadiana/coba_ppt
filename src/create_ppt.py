from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# ── Color Palette ──
BG_DARK    = RGBColor(0x1B, 0x1B, 0x2F)   # deep navy
BG_CARD    = RGBColor(0x27, 0x27, 0x44)   # card background
ACCENT     = RGBColor(0xE8, 0x4D, 0x4D)   # red accent (medical/emergency)
ACCENT2    = RGBColor(0xFF, 0xA5, 0x00)   # orange warning
GREEN      = RGBColor(0x2E, 0xCC, 0x71)   # green (do's)
RED        = RGBColor(0xE7, 0x4C, 0x3C)   # red (don'ts)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xBB, 0xBB, 0xCC)
GOLD       = RGBColor(0xFF, 0xD7, 0x00)   # golden hour
TEAL       = RGBColor(0x00, 0xB4, 0xD8)   # teal accent

def set_slide_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_shape_bg(slide, left, top, width, height, color, alpha=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape

def add_circle(slide, left, top, size, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def add_text_box(slide, left, top, width, height, text, font_size=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT, font_name='Calibri'):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = align
    return txBox

def add_rich_text_box(slide, left, top, width, height):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    return tf

# ══════════════════════════════════════════════════════════════
# SLIDE 1 — TITLE SLIDE
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
set_slide_bg(slide, BG_DARK)

# Decorative cross (medical symbol)
cross_v = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.15), Inches(0.8), Inches(0.5), Inches(1.5))
cross_v.fill.solid(); cross_v.fill.fore_color.rgb = ACCENT; cross_v.line.fill.background()
cross_h = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.55), Inches(1.25), Inches(1.7), Inches(0.5))
cross_h.fill.solid(); cross_h.fill.fore_color.rgb = ACCENT; cross_h.line.fill.background()

# Red accent bar top
bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.08))
bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()

add_text_box(slide, Inches(1), Inches(2.8), Inches(11.3), Inches(1.2),
             "PANDUAN LENGKAP", font_size=44, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text_box(slide, Inches(1), Inches(3.7), Inches(11.3), Inches(1),
             "PERTOLONGAN PERTAMA PADA KORBAN KECELAKAAN", font_size=30, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

# Divider line
divider = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(4.9), Inches(4.3), Inches(0.04))
divider.fill.solid(); divider.fill.fore_color.rgb = LIGHT_GRAY; divider.line.fill.background()

add_text_box(slide, Inches(1), Inches(5.1), Inches(11.3), Inches(0.6),
             "Berdasarkan Pedoman Basic Trauma Life Support (BTLS) & Prinsip DRSABCD",
             font_size=16, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)

add_text_box(slide, Inches(1), Inches(6.2), Inches(11.3), Inches(0.5),
             "Palang Merah Indonesia  |  Layanan Darurat: 119  •  112",
             font_size=14, color=TEAL, align=PP_ALIGN.CENTER)

# ══════════════════════════════════════════════════════════════
# SLIDE 2 — PENDAHULUAN
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_DARK)

bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.06))
bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()

# Section number
circle = add_circle(slide, Inches(0.6), Inches(0.4), Inches(0.7), ACCENT)
tf = circle.text_frame; tf.word_wrap = False
p = tf.paragraphs[0]; p.text = "00"; p.font.size = Pt(20); p.font.color.rgb = WHITE; p.font.bold = True; p.alignment = PP_ALIGN.CENTER
tf.paragraphs[0].space_before = Pt(6)

add_text_box(slide, Inches(1.5), Inches(0.45), Inches(5), Inches(0.6),
             "PENDAHULUAN", font_size=28, color=WHITE, bold=True)

# Content card
card = add_shape_bg(slide, Inches(0.6), Inches(1.5), Inches(12.1), Inches(5.5), BG_CARD)

tf = add_rich_text_box(slide, Inches(1.2), Inches(1.8), Inches(10.9), Inches(5))

p = tf.paragraphs[0]
p.text = "Mengapa Pertolongan Pertama Penting?"
p.font.size = Pt(24); p.font.color.rgb = ACCENT; p.font.bold = True

p = tf.add_paragraph(); p.space_before = Pt(20)
p.text = "Kecelakaan lalu lintas atau insiden tak terduga lainnya dapat terjadi kapan saja."
p.font.size = Pt(18); p.font.color.rgb = WHITE

p = tf.add_paragraph(); p.space_before = Pt(14)
p.text = "Penanganan pertama yang cepat dan tepat sangat krusial untuk:"
p.font.size = Pt(18); p.font.color.rgb = WHITE

bullets = [
    "Mencegah cedera yang lebih parah",
    "Menyelamatkan nyawa korban",
    "Memberikan waktu bagi tenaga medis profesional untuk tiba",
    "Mengurangi risiko komplikasi jangka panjang"
]
for b in bullets:
    p = tf.add_paragraph(); p.space_before = Pt(10); p.level = 1
    run = p.add_run(); run.text = "▸  "; run.font.size = Pt(18); run.font.color.rgb = ACCENT
    run2 = p.add_run(); run2.text = b; run2.font.size = Pt(18); run2.font.color.rgb = WHITE

p = tf.add_paragraph(); p.space_before = Pt(24)
run = p.add_run(); run.text = "Prinsip Utama: "; run.font.size = Pt(18); run.font.color.rgb = GOLD; run.font.bold = True
run2 = p.add_run(); run2.text = "DRSABCD (Danger, Response, Send for help, Airway, Breathing, Circulation, Defibrillation) + 3A (Aman Diri, Aman Lingkungan, Aman Korban)"
run2.font.size = Pt(17); run2.font.color.rgb = LIGHT_GRAY

# ══════════════════════════════════════════════════════════════
# SLIDE 3 — LANGKAH 1: AMANKAN KEADAAN (3A)
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_DARK)

bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.06))
bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()

circle = add_circle(slide, Inches(0.6), Inches(0.4), Inches(0.7), ACCENT)
tf = circle.text_frame
p = tf.paragraphs[0]; p.text = "01"; p.font.size = Pt(20); p.font.color.rgb = WHITE; p.font.bold = True; p.alignment = PP_ALIGN.CENTER
p.space_before = Pt(6)

add_text_box(slide, Inches(1.5), Inches(0.45), Inches(8), Inches(0.6),
             "AMANKAN KEADAAN — Prinsip 3A", font_size=28, color=WHITE, bold=True)

# Three cards for 3A
card_data = [
    ("🛡️", "AMAN DIRI", "Jangan menjadi korban selanjutnya. Pakai alat pelindung diri jika ada (seperti sarung tangan medis) dan pastikan posisi Anda aman dari lalu lintas.", ACCENT),
    ("🏛️", "AMAN LINGKUNGAN", "Matikan mesin kendaraan yang terlibat kecelakaan, pasang segitiga pengaman, atau minta orang lain untuk mengatur lalu lintas.", ACCENT2),
    ("🚑", "AMAN KORBAN", "Jauhkan korban dari bahaya yang mengancam nyawa secara langsung (seperti api atau kendaraan yang akan meledak).", TEAL),
]

for i, (icon, title, desc, color) in enumerate(card_data):
    left = Inches(0.6 + i * 4.1)
    card = add_shape_bg(slide, left, Inches(1.5), Inches(3.8), Inches(5.3), BG_CARD)

    # Color bar at top of card
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(1.5), Inches(3.8), Inches(0.06))
    top_bar.fill.solid(); top_bar.fill.fore_color.rgb = color; top_bar.line.fill.background()

    add_text_box(slide, left + Inches(0.3), Inches(1.8), Inches(3.2), Inches(0.8),
                 icon, font_size=40, align=PP_ALIGN.CENTER)

    add_text_box(slide, left + Inches(0.3), Inches(2.7), Inches(3.2), Inches(0.6),
                 title, font_size=22, color=color, bold=True, align=PP_ALIGN.CENTER)

    add_text_box(slide, left + Inches(0.3), Inches(3.5), Inches(3.2), Inches(3),
                 desc, font_size=16, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)

# ══════════════════════════════════════════════════════════════
# SLIDE 4 — LANGKAH 2: HUBUNGI BANTUAN
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_DARK)

bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.06))
bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()

circle = add_circle(slide, Inches(0.6), Inches(0.4), Inches(0.7), ACCENT)
tf = circle.text_frame
p = tf.paragraphs[0]; p.text = "02"; p.font.size = Pt(20); p.font.color.rgb = WHITE; p.font.bold = True; p.alignment = PP_ALIGN.CENTER
p.space_before = Pt(6)

add_text_box(slide, Inches(1.5), Inches(0.45), Inches(10), Inches(0.6),
             "HUBUNGI BANTUAN MEDIS DARURAT", font_size=28, color=WHITE, bold=True)

# Main card
card = add_shape_bg(slide, Inches(0.6), Inches(1.5), Inches(7.5), Inches(5.3), BG_CARD)

tf = add_rich_text_box(slide, Inches(1.2), Inches(1.8), Inches(6.3), Inches(4.8))
p = tf.paragraphs[0]
p.text = "Langkah terpenting adalah memanggil tenaga ahli."
p.font.size = Pt(20); p.font.color.rgb = WHITE

p = tf.add_paragraph(); p.space_before = Pt(20)
p.text = "Di Indonesia, Anda dapat menghubungi:"
p.font.size = Pt(18); p.font.color.rgb = LIGHT_GRAY

p = tf.add_paragraph(); p.space_before = Pt(16)
run = p.add_run(); run.text = "📞  119"; run.font.size = Pt(32); run.font.color.rgb = ACCENT; run.font.bold = True
run2 = p.add_run(); run2.text = "   Ambulans / Gawat Darurat Medis"; run2.font.size = Pt(18); run2.font.color.rgb = WHITE

p = tf.add_paragraph(); p.space_before = Pt(12)
run = p.add_run(); run.text = "📞  112"; run.font.size = Pt(32); run.font.color.rgb = ACCENT2; run.font.bold = True
run2 = p.add_run(); run2.text = "   Layanan Panggilan Darurat Terpadu"; run2.font.size = Pt(18); run2.font.color.rgb = WHITE

p = tf.add_paragraph(); p.space_before = Pt(24)
p.text = "Informasi yang perlu disampaikan:"
p.font.size = Pt(18); p.font.color.rgb = GOLD; p.font.bold = True

info_items = ["Lokasi kejadian (alamat jelas)", "Jumlah korban", "Kondisi korban secara singkat", "Jenis kecelakaan"]
for item in info_items:
    p = tf.add_paragraph(); p.space_before = Pt(8); p.level = 1
    run = p.add_run(); run.text = "▸  "; run.font.size = Pt(16); run.font.color.rgb = TEAL
    run2 = p.add_run(); run2.text = item; run2.font.size = Pt(16); run2.font.color.rgb = WHITE

# Side highlight card
highlight = add_shape_bg(slide, Inches(8.5), Inches(1.5), Inches(4.2), Inches(5.3), RGBColor(0x3A, 0x1A, 0x1A))

tf2 = add_rich_text_box(slide, Inches(8.9), Inches(1.8), Inches(3.4), Inches(4.8))
p = tf2.paragraphs[0]
p.text = "⚠️"; p.font.size = Pt(48); p.alignment = PP_ALIGN.CENTER

p = tf2.add_paragraph(); p.space_before = Pt(16)
p.text = "INGAT!"; p.font.size = Pt(24); p.font.color.rgb = ACCENT; p.font.bold = True; p.alignment = PP_ALIGN.CENTER

p = tf2.add_paragraph(); p.space_before = Pt(16)
p.text = "Jangan mencoba menolong dengan cara yang salah. Kecepatan menghubungi 119 jauh lebih berharga."
p.font.size = Pt(16); p.font.color.rgb = LIGHT_GRAY; p.alignment = PP_ALIGN.CENTER

p = tf2.add_paragraph(); p.space_before = Pt(20)
p.text = "\"The Golden Hour\""; p.font.size = Pt(18); p.font.color.rgb = GOLD; p.font.bold = True; p.alignment = PP_ALIGN.CENTER

p = tf2.add_paragraph(); p.space_before = Pt(8)
p.text = "60 menit pertama setelah trauma menentukan keselamatan korban!"; p.font.size = Pt(14); p.font.color.rgb = LIGHT_GRAY; p.alignment = PP_ALIGN.CENTER

# ══════════════════════════════════════════════════════════════
# SLIDE 5 — LANGKAH 3: PERIKSA KESADARAN
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_DARK)

bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.06))
bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()

circle = add_circle(slide, Inches(0.6), Inches(0.4), Inches(0.7), ACCENT)
tf = circle.text_frame
p = tf.paragraphs[0]; p.text = "03"; p.font.size = Pt(20); p.font.color.rgb = WHITE; p.font.bold = True; p.alignment = PP_ALIGN.CENTER
p.space_before = Pt(6)

add_text_box(slide, Inches(1.5), Inches(0.45), Inches(10), Inches(0.6),
             "PERIKSA KESADARAN KORBAN (Cek Respon)", font_size=28, color=WHITE, bold=True)

# Instruction card
card = add_shape_bg(slide, Inches(0.6), Inches(1.5), Inches(12.1), Inches(2.2), BG_CARD)
tf = add_rich_text_box(slide, Inches(1.2), Inches(1.7), Inches(10.9), Inches(1.8))
p = tf.paragraphs[0]
p.text = "Cara Memeriksa:"
p.font.size = Pt(20); p.font.color.rgb = TEAL; p.font.bold = True

p = tf.add_paragraph(); p.space_before = Pt(12)
p.text = "Dekati korban dan periksa kesadarannya dengan menepuk bahu secara perlahan sambil memanggil:"
p.font.size = Pt(17); p.font.color.rgb = WHITE

p = tf.add_paragraph(); p.space_before = Pt(12)
run = p.add_run(); run.text = "\"Pak/Bu, apakah Anda bisa mendengar saya?\""; run.font.size = Pt(20); run.font.color.rgb = GOLD; run.font.bold = True; run.font.italic = True

# Two scenario cards
# If conscious
card1 = add_shape_bg(slide, Inches(0.6), Inches(4.0), Inches(5.9), Inches(3), BG_CARD)
top_bar1 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(4.0), Inches(5.9), Inches(0.06))
top_bar1.fill.solid(); top_bar1.fill.fore_color.rgb = GREEN; top_bar1.line.fill.background()

tf1 = add_rich_text_box(slide, Inches(1.0), Inches(4.3), Inches(5.1), Inches(2.5))
p = tf1.paragraphs[0]
run = p.add_run(); run.text = "✅  "; run.font.size = Pt(22); run.font.color.rgb = GREEN
run2 = p.add_run(); run2.text = "JIKA SADAR"; run2.font.size = Pt(22); run2.font.color.rgb = GREEN; run2.font.bold = True

p = tf1.add_paragraph(); p.space_before = Pt(14)
p.text = "• Tenangkan korban"; p.font.size = Pt(17); p.font.color.rgb = WHITE
p = tf1.add_paragraph(); p.space_before = Pt(8)
p.text = "• Larang mereka untuk banyak bergerak"; p.font.size = Pt(17); p.font.color.rgb = WHITE
p = tf1.add_paragraph(); p.space_before = Pt(8)
p.text = "• Ajak mengobrol agar tidak panik"; p.font.size = Pt(17); p.font.color.rgb = WHITE

# If unconscious
card2 = add_shape_bg(slide, Inches(6.8), Inches(4.0), Inches(5.9), Inches(3), BG_CARD)
top_bar2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(4.0), Inches(5.9), Inches(0.06))
top_bar2.fill.solid(); top_bar2.fill.fore_color.rgb = RED; top_bar2.line.fill.background()

tf2 = add_rich_text_box(slide, Inches(7.2), Inches(4.3), Inches(5.1), Inches(2.5))
p = tf2.paragraphs[0]
run = p.add_run(); run.text = "❌  "; run.font.size = Pt(22); run.font.color.rgb = RED
run2 = p.add_run(); run2.text = "JIKA TIDAK SADAR"; run2.font.size = Pt(22); run2.font.color.rgb = RED; run2.font.bold = True

p = tf2.add_paragraph(); p.space_before = Pt(14)
p.text = "• Periksa pernapasan (lihat pergerakan dada)"; p.font.size = Pt(17); p.font.color.rgb = WHITE
p = tf2.add_paragraph(); p.space_before = Pt(8)
p.text = "• Jika tidak bernapas → lakukan RJP/CPR"; p.font.size = Pt(17); p.font.color.rgb = WHITE
p = tf2.add_paragraph(); p.space_before = Pt(8)
p.text = "• Kompresi dada (untuk yang terlatih)"; p.font.size = Pt(17); p.font.color.rgb = WHITE

# ══════════════════════════════════════════════════════════════
# SLIDE 6 — LANGKAH 4: JANGAN PINDAHKAN KORBAN
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_DARK)

bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.06))
bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()

circle = add_circle(slide, Inches(0.6), Inches(0.4), Inches(0.7), ACCENT)
tf = circle.text_frame
p = tf.paragraphs[0]; p.text = "04"; p.font.size = Pt(20); p.font.color.rgb = WHITE; p.font.bold = True; p.alignment = PP_ALIGN.CENTER
p.space_before = Pt(6)

add_text_box(slide, Inches(1.5), Inches(0.45), Inches(10), Inches(0.6),
             "JANGAN PINDAHKAN KORBAN SECARA SEMBARANGAN!", font_size=28, color=ACCENT, bold=True)

# Warning banner
warn_bg = add_shape_bg(slide, Inches(0.6), Inches(1.4), Inches(12.1), Inches(1.4), RGBColor(0x5C, 0x1A, 0x1A))
tf = add_rich_text_box(slide, Inches(1.2), Inches(1.5), Inches(10.9), Inches(1.2))
p = tf.paragraphs[0]
run = p.add_run(); run.text = "⚠️  KESALAHAN PALING FATAL YANG SERING TERJADI DI MASYARAKAT  ⚠️"
run.font.size = Pt(20); run.font.color.rgb = ACCENT; run.font.bold = True; p.alignment = PP_ALIGN.CENTER

# Golden rule card
card = add_shape_bg(slide, Inches(0.6), Inches(3.1), Inches(12.1), Inches(1.8), BG_CARD)
top_gold = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(3.1), Inches(12.1), Inches(0.06))
top_gold.fill.solid(); top_gold.fill.fore_color.rgb = GOLD; top_gold.line.fill.background()

tf = add_rich_text_box(slide, Inches(1.2), Inches(3.3), Inches(10.9), Inches(1.4))
p = tf.paragraphs[0]
run = p.add_run(); run.text = "🏆  ATURAN EMAS: "; run.font.size = Pt(20); run.font.color.rgb = GOLD; run.font.bold = True
run2 = p.add_run(); run2.text = "Kecuali ada ancaman nyawa langsung (seperti kebakaran atau risiko terlindas), JANGAN MEMINDAHKAN KORBAN."
run2.font.size = Pt(18); run2.font.color.rgb = WHITE

# Consequence card
card2 = add_shape_bg(slide, Inches(0.6), Inches(5.2), Inches(12.1), Inches(2), RGBColor(0x3A, 0x1A, 0x1A))
tf2 = add_rich_text_box(slide, Inches(1.2), Inches(5.4), Inches(10.9), Inches(1.6))
p = tf2.paragraphs[0]
run = p.add_run(); run.text = "Konsekuensi Memindahkan Korban Tanpa Teknik yang Benar:"; run.font.size = Pt(18); run.font.color.rgb = RED; run.font.bold = True

p = tf2.add_paragraph(); p.space_before = Pt(12)
p.text = "▸  Kerusakan tulang belakang dan leher"; p.font.size = Pt(17); p.font.color.rgb = WHITE
p = tf2.add_paragraph(); p.space_before = Pt(8)
p.text = "▸  Kelumpuhan permanen (tetap lumpuh seumur hidup)"; p.font.size = Pt(17); p.font.color.rgb = WHITE
p = tf2.add_paragraph(); p.space_before = Pt(8)
p.text = "▸  Kematian akibat cedera yang memburuk"; p.font.size = Pt(17); p.font.color.rgb = WHITE

# ══════════════════════════════════════════════════════════════
# SLIDE 7 — LANGKAH 5: TANGANI PERDARAHAN
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_DARK)

bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.06))
bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()

circle = add_circle(slide, Inches(0.6), Inches(0.4), Inches(0.7), ACCENT)
tf = circle.text_frame
p = tf.paragraphs[0]; p.text = "05"; p.font.size = Pt(20); p.font.color.rgb = WHITE; p.font.bold = True; p.alignment = PP_ALIGN.CENTER
p.space_before = Pt(6)

add_text_box(slide, Inches(1.5), Inches(0.45), Inches(10), Inches(0.6),
             "TANGANI PERDARAHAN HEBAT", font_size=28, color=WHITE, bold=True)

# Main instruction card
card = add_shape_bg(slide, Inches(0.6), Inches(1.5), Inches(12.1), Inches(1.8), BG_CARD)
tf = add_rich_text_box(slide, Inches(1.2), Inches(1.7), Inches(10.9), Inches(1.4))
p = tf.paragraphs[0]
p.text = "Jika Anda melihat perdarahan yang menyemprot atau mengalir deras:"
p.font.size = Pt(20); p.font.color.rgb = WHITE

p = tf.add_paragraph(); p.space_before = Pt(10)
p.text = "Ini adalah situasi darurat yang memerlukan tindakan cepat!"
p.font.size = Pt(17); p.font.color.rgb = ACCENT; p.font.bold = True

# Steps cards
steps = [
    ("1", "TEKAN LANGSUNG", "Gunakan kain bersih (baju, handuk, atau kassa) untuk menekan langsung pada titik luka.", TEAL),
    ("2", "PERTAHANKAN TEKANAN", "Pertahankan tekanan secara konsisten sampai bantuan medis tiba. Jangan mengangkat untuk mengecek.", ACCENT2),
    ("3", "PASTIKAN TEKANAN KUAT", "Tekanan yang kuat dan konsisten adalah kunci menghentikan perdarahan hebat.", ACCENT),
]

for i, (num, title, desc, color) in enumerate(steps):
    left = Inches(0.6 + i * 4.1)
    card = add_shape_bg(slide, left, Inches(3.6), Inches(3.8), Inches(3.5), BG_CARD)
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(3.6), Inches(3.8), Inches(0.06))
    top_bar.fill.solid(); top_bar.fill.fore_color.rgb = color; top_bar.line.fill.background()

    circle_num = add_circle(slide, left + Inches(1.4), Inches(3.9), Inches(1), color)
    tf_num = circle_num.text_frame
    p = tf_num.paragraphs[0]; p.text = num; p.font.size = Pt(32); p.font.color.rgb = WHITE; p.font.bold = True; p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(4)

    add_text_box(slide, left + Inches(0.3), Inches(5.1), Inches(3.2), Inches(0.5),
                 title, font_size=18, color=color, bold=True, align=PP_ALIGN.CENTER)

    add_text_box(slide, left + Inches(0.3), Inches(5.7), Inches(3.2), Inches(1.2),
                 desc, font_size=15, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)

# ══════════════════════════════════════════════════════════════
# SLIDE 8 — DO'S AND DON'TS
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_DARK)

bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.06))
bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()

circle = add_circle(slide, Inches(0.6), Inches(0.4), Inches(0.7), ACCENT)
tf = circle.text_frame
p = tf.paragraphs[0]; p.text = "06"; p.font.size = Pt(20); p.font.color.rgb = WHITE; p.font.bold = True; p.alignment = PP_ALIGN.CENTER
p.space_before = Pt(6)

add_text_box(slide, Inches(1.5), Inches(0.45), Inches(10), Inches(0.6),
             "DO'S & DON'TS — Yang Boleh dan Tidak Boleh Dilakukan", font_size=28, color=WHITE, bold=True)

# DO's column
do_card = add_shape_bg(slide, Inches(0.6), Inches(1.4), Inches(5.9), Inches(5.6), BG_CARD)
do_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.4), Inches(5.9), Inches(0.08))
do_bar.fill.solid(); do_bar.fill.fore_color.rgb = GREEN; do_bar.line.fill.background()

add_text_box(slide, Inches(0.6), Inches(1.6), Inches(5.9), Inches(0.6),
             "✅  BOLEH DILAKUKAN (Do's)", font_size=22, color=GREEN, bold=True, align=PP_ALIGN.CENTER)

dos = [
    ("Menjaga suhu tubuh korban", "Agar tetap hangat menggunakan jaket atau selimut."),
    ("Membuka kaca helm pelindung", "Untuk melancarkan udara (jika korban memakai helm full-face)."),
    ("Mengajak korban mengobrol", "Jika sadar, agar mereka tidak panik atau tertidur."),
]

for i, (title, desc) in enumerate(dos):
    y = Inches(2.4 + i * 1.5)
    tf = add_rich_text_box(slide, Inches(1.0), y, Inches(5.1), Inches(1.3))
    p = tf.paragraphs[0]
    run = p.add_run(); run.text = "▸  "; run.font.size = Pt(18); run.font.color.rgb = GREEN
    run2 = p.add_run(); run2.text = title; run2.font.size = Pt(18); run2.font.color.rgb = WHITE; run2.font.bold = True

    p = tf.add_paragraph(); p.space_before = Pt(4)
    p.text = desc; p.font.size = Pt(15); p.font.color.rgb = LIGHT_GRAY

# DON'Ts column
dont_card = add_shape_bg(slide, Inches(6.8), Inches(1.4), Inches(5.9), Inches(5.6), BG_CARD)
dont_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(1.4), Inches(5.9), Inches(0.08))
dont_bar.fill.solid(); dont_bar.fill.fore_color.rgb = RED; dont_bar.line.fill.background()

add_text_box(slide, Inches(6.8), Inches(1.6), Inches(5.9), Inches(0.6),
             "❌  DILARANG KERAS (Don'ts)", font_size=22, color=RED, bold=True, align=PP_ALIGN.CENTER)

donts = [
    ("Memberikan makan atau minum", "Berisiko tersedak dan menutup jalan napas korban."),
    ("Melepas helm secara paksa", "Dari kepala korban, karena berisiko merusak tulang leher."),
    ("Mengerumuni korban terlalu rapat", "Menghalangi sirkulasi udara segar untuk korban."),
]

for i, (title, desc) in enumerate(donts):
    y = Inches(2.4 + i * 1.5)
    tf = add_rich_text_box(slide, Inches(7.2), y, Inches(5.1), Inches(1.3))
    p = tf.paragraphs[0]
    run = p.add_run(); run.text = "▸  "; run.font.size = Pt(18); run.font.color.rgb = RED
    run2 = p.add_run(); run2.text = title; run2.font.size = Pt(18); run2.font.color.rgb = WHITE; run2.font.bold = True

    p = tf.add_paragraph(); p.space_before = Pt(4)
    p.text = desc; p.font.size = Pt(15); p.font.color.rgb = LIGHT_GRAY

# ══════════════════════════════════════════════════════════════
# SLIDE 9 — THE GOLDEN HOUR
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_DARK)

bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.06))
bar.fill.solid(); bar.fill.fore_color.rgb = GOLD; bar.line.fill.background()

# Big golden hour visual
gold_bg = add_shape_bg(slide, Inches(2), Inches(0.8), Inches(9.3), Inches(6), RGBColor(0x2A, 0x25, 0x10))
gold_border = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2), Inches(0.8), Inches(9.3), Inches(6))
gold_border.fill.background(); gold_border.line.color.rgb = GOLD; gold_border.line.width = Pt(2)

# Clock icon area
clock_circle = add_circle(slide, Inches(5.65), Inches(1.2), Inches(2), GOLD)
tf_clock = clock_circle.text_frame
p = tf_clock.paragraphs[0]; p.text = "⏱️"; p.font.size = Pt(48); p.alignment = PP_ALIGN.CENTER; p.space_before = Pt(12)

add_text_box(slide, Inches(2.5), Inches(3.4), Inches(8.3), Inches(0.7),
             "THE GOLDEN HOUR", font_size=36, color=GOLD, bold=True, align=PP_ALIGN.CENTER)

add_text_box(slide, Inches(2.5), Inches(4.1), Inches(8.3), Inches(0.5),
             "\"Jam Emas\" — 60 Menit Pertama Setelah Kecelakaan", font_size=18, color=WHITE, align=PP_ALIGN.CENTER)

divider = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(4.8), Inches(4.3), Inches(0.03))
divider.fill.solid(); divider.fill.fore_color.rgb = GOLD; divider.line.fill.background()

tf = add_rich_text_box(slide, Inches(3), Inches(5.1), Inches(7.3), Inches(1.5))
p = tf.paragraphs[0]
p.text = "Penanganan yang tepat dalam waktu ini sangat menentukan tingkat keselamatan dan pemulihan korban."
p.font.size = Pt(17); p.font.color.rgb = WHITE; p.alignment = PP_ALIGN.CENTER

p = tf.add_paragraph(); p.space_before = Pt(14)
run = p.add_run(); run.text = "Kecepatan menghubungi 119 jauh lebih berharga "; run.font.size = Pt(17); run.font.color.rgb = WHITE
run2 = p.add_run(); run2.text = "daripada mencoba menolong dengan cara yang salah."; run2.font.size = Pt(17); run2.font.color.rgb = ACCENT; run2.font.bold = True
p.alignment = PP_ALIGN.CENTER

# ══════════════════════════════════════════════════════════════
# SLIDE 10 — RINGKASAN & PENUTUP
# ══════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BG_DARK)

bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.06))
bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()

add_text_box(slide, Inches(1), Inches(0.5), Inches(11.3), Inches(0.8),
             "RINGKASAN — 5 Langkah Pertolongan Pertama", font_size=28, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

# Summary cards in a row
summary = [
    ("01", "AMANKAN\nKEADAAN", "Prinsip 3A:\nDiri, Lingkungan,\nKorban", ACCENT),
    ("02", "HUBUNGI\nBANTUAN", "Dial 119 / 112\nAmbulans &\nGawat Darurat", ACCENT2),
    ("03", "PERIKSA\nKESADARAN", "Cek Respon:\nTepuk bahu,\npanggil korban", TEAL),
    ("04", "JANGAN\nPINDAHKAN", "Aturan Emas:\nKecuali ancaman\nlangsung", GOLD),
    ("05", "TANGANI\nPERDARAHAN", "Tekan luka\ndengan kain\nbersih & kuat", GREEN),
]

for i, (num, title, desc, color) in enumerate(summary):
    left = Inches(0.4 + i * 2.55)
    card = add_shape_bg(slide, left, Inches(1.6), Inches(2.35), Inches(5.2), BG_CARD)
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(1.6), Inches(2.35), Inches(0.06))
    top_bar.fill.solid(); top_bar.fill.fore_color.rgb = color; top_bar.line.fill.background()

    circle_num = add_circle(slide, left + Inches(0.65), Inches(1.9), Inches(1), color)
    tf_num = circle_num.text_frame
    p = tf_num.paragraphs[0]; p.text = num; p.font.size = Pt(28); p.font.color.rgb = WHITE; p.font.bold = True; p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(4)

    add_text_box(slide, left + Inches(0.15), Inches(3.1), Inches(2.05), Inches(1),
                 title, font_size=16, color=color, bold=True, align=PP_ALIGN.CENTER)

    add_text_box(slide, left + Inches(0.15), Inches(4.2), Inches(2.05), Inches(2.2),
                 desc, font_size=13, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)

# Bottom message
add_text_box(slide, Inches(1), Inches(7.0), Inches(11.3), Inches(0.4),
             "📞  119  •  112  |  Selamatkan Nyawa dengan Penanganan yang Tepat!",
             font_size=16, color=TEAL, bold=True, align=PP_ALIGN.CENTER)

# ══════════════════════════════════════════════════════════════
# SAVE
# ══════════════════════════════════════════════════════════════
output_path = "/home/runner/coba_ppt/Panduan_Pertolongan_Pertama_Korban_Kecelakaan.pptx"
prs.save(output_path)
print(f"✅ PPT saved to: {output_path}")
